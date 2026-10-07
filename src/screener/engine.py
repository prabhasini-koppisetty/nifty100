"""
Screener Engine Module for Nifty 100 Financial Intelligence Platform.
Provides screening algorithms, preset execution, threshold filtering,
composite quality scoring with P10/P90 winsorisation, and Excel report generation.
"""

import sys
import sqlite3
import yaml
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parents[2]
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

import pandas as pd
import numpy as np
from typing import Dict, List, Any, Optional, Tuple

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT_DIR = Path(__file__).resolve().parents[2]
DB_PATH = ROOT_DIR / "db" / "nifty100.db"
CONFIG_PATH = ROOT_DIR / "config" / "screener_config.yaml"
OUTPUT_DIR = ROOT_DIR / "output"
OUTPUT_DIR.mkdir(exist_ok=True)
SCREENER_EXCEL_PATH = OUTPUT_DIR / "screener_output.xlsx"


def load_screener_config(config_path: Optional[Path] = None) -> Dict[str, Any]:
    """Load screener configuration from YAML file."""
    path = config_path or CONFIG_PATH
    if not path.exists():
        raise FileNotFoundError(f"Screener configuration file not found at {path}")
    with open(path, "r", encoding="utf-8") as f:
        return yaml.safe_load(f)


def calculate_sprint3_composite_score(df: pd.DataFrame) -> pd.DataFrame:
    """
    Calculate 100-point Composite Quality Score with P10/P90 Winsorisation.
    
    Pillars & Weights:
    1. Profitability (35%):
       - ROE: 15%
       - ROCE: 10%
       - NPM: 10%
    2. Cash Quality (30%):
       - FCF Growth / Value: 15%
       - CFO / PAT ratio: 10%
       - FCF positive flag: 5%
    3. Growth (20%):
       - Revenue CAGR 5yr: 10%
       - PAT CAGR 5yr: 10%
    4. Leverage (15%):
       - D/E score (lower is better, inverted): 10%
       - ICR score: 5%
       
    Includes sector-relative score calculated within broad_sector.
    """
    res = df.copy()
    
    # Required raw metric mappings
    metrics_spec = {
        "roe": ("return_on_equity_pct", 15.0, False),
        "roce": ("return_on_capital_employed_pct", 10.0, False),
        "npm": ("net_profit_margin_pct", 10.0, False),
        "fcf": ("free_cash_flow_cr", 15.0, False),
        "cfo_pat": ("cfo_quality_score", 10.0, False),
        "rev_cagr_5yr": ("revenue_cagr_5yr", 10.0, False),
        "pat_cagr_5yr": ("pat_cagr_5yr", 10.0, False),
        "de": ("debt_to_equity", 10.0, True),  # Inverted: lower is better
        "icr": ("interest_coverage", 5.0, False),
    }

    def winsorise_and_normalize(series: pd.Series, invert: bool = False) -> pd.Series:
        valid = series.dropna()
        if len(valid) == 0:
            return pd.Series(np.nan, index=series.index)
        p10 = valid.quantile(0.10)
        p90 = valid.quantile(0.90)
        
        # Clip
        clipped = series.clip(lower=p10, upper=p90)
        
        if p90 == p10:
            norm = pd.Series(50.0, index=series.index)
        else:
            norm = (clipped - p10) / (p90 - p10) * 100.0
            
        if invert:
            norm = 100.0 - norm
        return norm

    # Calculate overall normalized component scores
    weighted_scores = []
    total_weights = []

    for name, (col, weight, invert) in metrics_spec.items():
        if col in res.columns:
            norm_col = winsorise_and_normalize(res[col], invert=invert)
            res[f"norm_{name}"] = norm_col
            weighted_scores.append(norm_col.fillna(0) * weight)
            total_weights.append((~norm_col.isna()).astype(float) * weight)

    # Special handling for FCF positive flag (5%)
    if "free_cash_flow_cr" in res.columns:
        fcf_pos = (res["free_cash_flow_cr"] > 0).astype(float) * 100.0
        res["norm_fcf_pos"] = fcf_pos
        weighted_scores.append(fcf_pos * 5.0)
        total_weights.append(pd.Series(5.0, index=res.index))

    sum_weighted = pd.DataFrame(weighted_scores).sum(axis=0)
    sum_weights = pd.DataFrame(total_weights).sum(axis=0)

    # Avoid division by zero
    composite_score = np.where(sum_weights > 0, (sum_weighted / sum_weights), np.nan)
    res["composite_quality_score"] = np.round(composite_score, 2)

    # Calculate sector-relative composite score
    sector_scores = []
    for sector, group in res.groupby("broad_sector"):
        grp_weighted = []
        grp_weights = []
        for name, (col, weight, invert) in metrics_spec.items():
            if col in group.columns:
                norm_col = winsorise_and_normalize(group[col], invert=invert)
                grp_weighted.append(norm_col.fillna(0) * weight)
                grp_weights.append((~norm_col.isna()).astype(float) * weight)
        if "free_cash_flow_cr" in group.columns:
            fcf_pos = (group["free_cash_flow_cr"] > 0).astype(float) * 100.0
            grp_weighted.append(fcf_pos * 5.0)
            grp_weights.append(pd.Series(5.0, index=group.index))
            
        s_weighted = pd.DataFrame(grp_weighted).sum(axis=0)
        s_weights = pd.DataFrame(grp_weights).sum(axis=0)
        s_score = np.where(s_weights > 0, (s_weighted / s_weights), np.nan)
        group_res = pd.Series(np.round(s_score, 2), index=group.index)
        sector_scores.append(group_res)

    if sector_scores:
        res["sector_composite_quality_score"] = pd.concat(sector_scores).reindex(res.index)
    else:
        res["sector_composite_quality_score"] = res["composite_quality_score"]

    return res


def load_financial_data(db_path: Optional[Path] = None) -> pd.DataFrame:
    """
    Load joined financial metrics for the latest annual period per company from SQLite.
    Includes financial ratios, market cap ratios, profit & loss, sector data, and company info.
    Includes all 92 Nifty 100 companies.
    """
    path = db_path or DB_PATH
    conn = sqlite3.connect(path)
    
    # Query to fetch latest annual ratios and joined data for all 92 companies
    query = """
    WITH latest_fr AS (
        SELECT fr.*,
               ROW_NUMBER() OVER (
                   PARTITION BY company_id 
                   ORDER BY (return_on_equity_pct IS NOT NULL) DESC, year DESC
               ) as rk
        FROM financial_ratios fr
    ),
    prev_de AS (
        SELECT company_id, year, debt_to_equity as prev_debt_to_equity,
               ROW_NUMBER() OVER (PARTITION BY company_id ORDER BY year DESC) as rk_prev
        FROM financial_ratios
        WHERE debt_to_equity IS NOT NULL
    )
    SELECT 
        c.id as company_id,
        c.company_name,
        s.broad_sector,
        s.sub_sector,
        s.market_cap_category,
        COALESCE(r.year, '2024-03') as year,
        r.return_on_equity_pct,
        r.return_on_capital_employed_pct,
        r.operating_profit_margin_pct,
        r.net_profit_margin_pct,
        r.debt_to_equity,
        r.interest_coverage,
        r.icr_label,
        r.free_cash_flow_cr,
        r.revenue_cagr_3yr,
        r.revenue_cagr_5yr,
        r.pat_cagr_5yr,
        r.eps_cagr_5yr,
        r.asset_turnover,
        r.dividend_payout_ratio_pct,
        r.cfo_quality_score,
        r.composite_quality_score as sprint2_composite_quality_score,
        mc.pe_ratio,
        mc.pb_ratio,
        mc.dividend_yield_pct,
        mc.market_cap_crore,
        pl.sales,
        pl.net_profit,
        pde.prev_debt_to_equity
    FROM companies c
    LEFT JOIN latest_fr r ON c.id = r.company_id AND r.rk = 1
    LEFT JOIN sectors s ON c.id = s.company_id
    LEFT JOIN market_cap mc ON c.id = mc.company_id AND (r.year = mc.year OR mc.year = '2024-03')
    LEFT JOIN profitandloss pl ON c.id = pl.company_id AND r.year = pl.year
    LEFT JOIN prev_de pde ON c.id = pde.company_id AND pde.rk_prev = 2
    ORDER BY c.id ASC
    """
    
    df = pd.read_sql(query, conn)
    conn.close()

    # Deduplicate in case JOIN created duplicates
    df = df.drop_duplicates(subset=["company_id"]).reset_index(drop=True)

    # Compute declining D/E YoY flag
    df["de_declining_yoy"] = False
    valid_de_mask = (~df["debt_to_equity"].isna()) & (~df["prev_debt_to_equity"].isna())
    df.loc[valid_de_mask, "de_declining_yoy"] = df.loc[valid_de_mask, "debt_to_equity"] < df.loc[valid_de_mask, "prev_debt_to_equity"]

    # Calculate Sprint 3 composite quality score
    df = calculate_sprint3_composite_score(df)

    return df


def apply_filters(df: pd.DataFrame, filters: Dict[str, Any]) -> pd.DataFrame:
    """
    Apply filter thresholds on financial DataFrame.
    
    Special Rules:
    1. Financials D/E Carve-out: Companies with broad_sector == 'Financials' skip D/E restriction.
    2. Debt Free ICR Rule: If icr_label == 'Debt Free', treat ICR as infinity (always pass ICR min threshold).
    3. Safe missing values: Null/NaN values fail filters unless explicitly allowed.
    """
    filtered = df.copy()

    # 1. ROE minimum
    roe_min = filters.get("roe_min") or filters.get("return_on_equity_pct_min")
    if roe_min is not None:
        filtered = filtered[filtered["return_on_equity_pct"].notna() & (filtered["return_on_equity_pct"] >= float(roe_min))]

    # 2. D/E maximum (Special rule for Financials)
    de_max = filters.get("de_max") or filters.get("debt_to_equity_max")
    if de_max is not None:
        de_threshold = float(de_max)
        is_financial = filtered["broad_sector"] == "Financials"
        de_pass = filtered["debt_to_equity"].notna() & (filtered["debt_to_equity"] <= de_threshold)
        # Financials automatically skip D/E restriction
        filtered = filtered[is_financial | de_pass]

    # 3. FCF minimum
    fcf_min = filters.get("fcf_min") or filters.get("free_cash_flow_cr_min")
    if fcf_min is not None:
        filtered = filtered[filtered["free_cash_flow_cr"].notna() & (filtered["free_cash_flow_cr"] > float(fcf_min))]

    # 4. Revenue CAGR 5yr minimum
    rev_cagr_5yr_min = filters.get("revenue_cagr_5yr_min")
    if rev_cagr_5yr_min is not None:
        filtered = filtered[filtered["revenue_cagr_5yr"].notna() & (filtered["revenue_cagr_5yr"] >= float(rev_cagr_5yr_min))]

    # 5. PAT CAGR 5yr minimum
    pat_cagr_5yr_min = filters.get("pat_cagr_5yr_min")
    if pat_cagr_5yr_min is not None:
        filtered = filtered[filtered["pat_cagr_5yr"].notna() & (filtered["pat_cagr_5yr"] >= float(pat_cagr_5yr_min))]

    # 6. OPM minimum
    opm_min = filters.get("opm_min") or filters.get("operating_profit_margin_pct_min")
    if opm_min is not None:
        filtered = filtered[filtered["operating_profit_margin_pct"].notna() & (filtered["operating_profit_margin_pct"] >= float(opm_min))]

    # 7. P/E maximum
    pe_max = filters.get("pe_max") or filters.get("pe_ratio_max")
    if pe_max is not None:
        filtered = filtered[filtered["pe_ratio"].notna() & (filtered["pe_ratio"] <= float(pe_max))]

    # 8. P/B maximum
    pb_max = filters.get("pb_max") or filters.get("pb_ratio_max")
    if pb_max is not None:
        filtered = filtered[filtered["pb_ratio"].notna() & (filtered["pb_ratio"] <= float(pb_max))]

    # 9. Dividend Yield minimum
    div_yield_min = filters.get("dividend_yield_min") or filters.get("dividend_yield_pct_min")
    if div_yield_min is not None:
        filtered = filtered[filtered["dividend_yield_pct"].notna() & (filtered["dividend_yield_pct"] >= float(div_yield_min))]

    # 10. ICR minimum (Special rule for Debt Free)
    icr_min = filters.get("icr_min") or filters.get("interest_coverage_min")
    if icr_min is not None:
        icr_threshold = float(icr_min)
        is_debt_free = filtered["icr_label"] == "Debt Free"
        icr_pass = filtered["interest_coverage"].notna() & (filtered["interest_coverage"] >= icr_threshold)
        # Debt free companies always pass ICR threshold
        filtered = filtered[is_debt_free | icr_pass]

    # 11. Market Cap minimum
    mcap_min = filters.get("market_cap_min") or filters.get("market_cap_crore_min")
    if mcap_min is not None:
        filtered = filtered[filtered["market_cap_crore"].notna() & (filtered["market_cap_crore"] >= float(mcap_min))]

    # 12. Net Profit minimum
    net_prof_min = filters.get("net_profit_min")
    if net_prof_min is not None:
        filtered = filtered[filtered["net_profit"].notna() & (filtered["net_profit"] >= float(net_prof_min))]

    # 13. EPS CAGR minimum
    eps_cagr_min = filters.get("eps_cagr_min") or filters.get("eps_cagr_5yr_min")
    if eps_cagr_min is not None:
        filtered = filtered[filtered["eps_cagr_5yr"].notna() & (filtered["eps_cagr_5yr"] >= float(eps_cagr_min))]

    # 14. Asset Turnover minimum
    asset_turnover_min = filters.get("asset_turnover_min")
    if asset_turnover_min is not None:
        filtered = filtered[filtered["asset_turnover"].notna() & (filtered["asset_turnover"] >= float(asset_turnover_min))]

    # 15. Sales minimum
    sales_min = filters.get("sales_min")
    if sales_min is not None:
        filtered = filtered[filtered["sales"].notna() & (filtered["sales"] >= float(sales_min))]

    # Additional preset criteria
    div_payout_max = filters.get("dividend_payout_ratio_pct_max")
    if div_payout_max is not None:
        filtered = filtered[filtered["dividend_payout_ratio_pct"].notna() & (filtered["dividend_payout_ratio_pct"] <= float(div_payout_max))]

    rev_cagr_3yr_min = filters.get("revenue_cagr_3yr_min")
    if rev_cagr_3yr_min is not None:
        filtered = filtered[filtered["revenue_cagr_3yr"].notna() & (filtered["revenue_cagr_3yr"] >= float(rev_cagr_3yr_min))]

    if filters.get("fcf_positive_latest"):
        filtered = filtered[filtered["free_cash_flow_cr"].notna() & (filtered["free_cash_flow_cr"] > 0)]

    if filters.get("de_declining_yoy"):
        filtered = filtered[filtered["de_declining_yoy"] == True]

    return filtered


def run_preset(df: pd.DataFrame, preset_key: str, config: Optional[Dict[str, Any]] = None) -> pd.DataFrame:
    """Run one of the 6 preset screeners and return results sorted by composite_quality_score DESC."""
    cfg = config or load_screener_config()
    presets = cfg.get("presets", {})
    
    if preset_key not in presets:
        raise ValueError(f"Unknown preset key: '{preset_key}'. Valid presets: {list(presets.keys())}")
        
    preset_def = presets[preset_key]
    criteria = preset_def.get("criteria", {})
    
    result = apply_filters(df, criteria)
    result = result.sort_values(by="composite_quality_score", ascending=False).reset_index(drop=True)
    return result


def run_custom_screener(df: pd.DataFrame, filters: Dict[str, Any]) -> pd.DataFrame:
    """Run a custom threshold screener with arbitrary criteria."""
    result = apply_filters(df, filters)
    result = result.sort_values(by="composite_quality_score", ascending=False).reset_index(drop=True)
    return result


def validate_screener_result(preset_name: str, df_result: pd.DataFrame) -> Dict[str, Any]:
    """Validate screener result metrics and record counts."""
    count = len(df_result)
    top_5 = df_result.head(5)[["company_id", "company_name", "composite_quality_score"]].to_dict("records") if count > 0 else []
    return {
        "preset_name": preset_name,
        "count": count,
        "in_expected_range": 5 <= count <= 50,
        "top_5": top_5
    }


def export_screener_results(preset_results: Dict[str, pd.DataFrame], output_path: Optional[Path] = None) -> Path:
    """
    Generate output/screener_output.xlsx containing exactly 6 sheets.
    Applies openpyxl formatting:
    - Headers: Dark blue fill, white bold font, freeze top row, autofilter.
    - Data rows: Conditional formatting (Green if meeting preset criteria, Red if failing).
    - Auto-adjust column widths.
    """
    out_file = output_path or SCREENER_EXCEL_PATH
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    pass_fill = PatternFill(start_color="E6F4EA", end_color="E6F4EA", fill_type="solid")
    fail_fill = PatternFill(start_color="FCE8E6", end_color="FCE8E6", fill_type="solid")
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0")
    )

    sheet_names = {
        "quality_compounder": "Quality Compounder",
        "value_pick": "Value Pick",
        "growth_accelerator": "Growth Accelerator",
        "dividend_champion": "Dividend Champion",
        "debt_free_blue_chip": "Debt-Free Blue Chip",
        "turnaround_watch": "Turnaround Watch"
    }

    # Preset criteria specs for cell-level color highlighting
    preset_threshold_checks = {
        "quality_compounder": {
            "return_on_equity_pct": lambda x: x >= 15.0 if pd.notna(x) else False,
            "debt_to_equity": lambda x: x <= 1.0 if pd.notna(x) else False,
            "free_cash_flow_cr": lambda x: x > 0 if pd.notna(x) else False,
            "revenue_cagr_5yr": lambda x: x >= 10.0 if pd.notna(x) else False,
        },
        "value_pick": {
            "pe_ratio": lambda x: x <= 20.0 if pd.notna(x) else False,
            "pb_ratio": lambda x: x <= 3.0 if pd.notna(x) else False,
            "debt_to_equity": lambda x: x <= 2.0 if pd.notna(x) else False,
            "dividend_yield_pct": lambda x: x >= 1.0 if pd.notna(x) else False,
        },
        "growth_accelerator": {
            "pat_cagr_5yr": lambda x: x >= 20.0 if pd.notna(x) else False,
            "revenue_cagr_5yr": lambda x: x >= 15.0 if pd.notna(x) else False,
            "debt_to_equity": lambda x: x <= 2.0 if pd.notna(x) else False,
        },
        "dividend_champion": {
            "dividend_yield_pct": lambda x: x >= 2.0 if pd.notna(x) else False,
            "dividend_payout_ratio_pct": lambda x: x <= 80.0 if pd.notna(x) else False,
            "free_cash_flow_cr": lambda x: x > 0 if pd.notna(x) else False,
        },
        "debt_free_blue_chip": {
            "debt_to_equity": lambda x: x <= 0.0 if pd.notna(x) else False,
            "return_on_equity_pct": lambda x: x >= 12.0 if pd.notna(x) else False,
            "sales": lambda x: x >= 5000.0 if pd.notna(x) else False,
        },
        "turnaround_watch": {
            "revenue_cagr_3yr": lambda x: x >= 10.0 if pd.notna(x) else False,
            "free_cash_flow_cr": lambda x: x > 0 if pd.notna(x) else False,
            "de_declining_yoy": lambda x: x == True if pd.notna(x) else False,
        }
    }

    cols_to_export = [
        "company_id", "company_name", "broad_sector", "sub_sector", "year",
        "composite_quality_score", "sector_composite_quality_score",
        "return_on_equity_pct", "return_on_capital_employed_pct", "operating_profit_margin_pct",
        "net_profit_margin_pct", "debt_to_equity", "interest_coverage", "free_cash_flow_cr",
        "revenue_cagr_3yr", "revenue_cagr_5yr", "pat_cagr_5yr", "eps_cagr_5yr",
        "pe_ratio", "pb_ratio", "dividend_yield_pct", "sales", "net_profit"
    ]

    for p_key, title in sheet_names.items():
        ws = wb.create_sheet(title=title)
        df_preset = preset_results.get(p_key, pd.DataFrame())
        
        # Ensure available columns
        exp_cols = [c for c in cols_to_export if c in df_preset.columns]
        
        # Write headers
        ws.append(exp_cols)
        for col_idx in range(1, len(exp_cols) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(exp_cols))}{len(df_preset) + 1}"

        # Write data rows
        checks = preset_threshold_checks.get(p_key, {})
        for row_idx, row_data in enumerate(df_preset[exp_cols].to_dict("records"), start=2):
            for col_idx, col_name in enumerate(exp_cols, start=1):
                val = row_data.get(col_name)
                # Convert numpy types to native Python
                if pd.isna(val):
                    cell_val = None
                elif isinstance(val, (np.integer, np.int64)):
                    cell_val = int(val)
                elif isinstance(val, (np.floating, np.float64)):
                    cell_val = round(float(val), 2)
                else:
                    cell_val = str(val)

                cell = ws.cell(row=row_idx, column=col_idx, value=cell_val)
                cell.font = regular_font
                cell.border = thin_border

                # Check if this metric is part of the preset criteria and color code
                if col_name in checks:
                    is_pass = checks[col_name](val)
                    cell.fill = pass_fill if is_pass else fail_fill
                elif col_name in ["company_id", "company_name", "broad_sector"]:
                    cell.alignment = Alignment(horizontal="left")
                else:
                    cell.alignment = Alignment(horizontal="right")

        # Auto-adjust column width
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    wb.save(out_file)
    return out_file


if __name__ == "__main__":
    print("Testing Screener Engine...")
    cfg = load_screener_config()
    df_data = load_financial_data()
    print(f"Loaded {len(df_data)} companies for screening.")

    results = {}
    for p_key in cfg.get("presets", {}):
        df_p = run_preset(df_data, p_key, cfg)
        results[p_key] = df_p
        val = validate_screener_result(p_key, df_p)
        print(f"Preset '{p_key}': {val['count']} companies matching criteria.")

    out_file = export_screener_results(results)
    print(f"Exported screener output to {out_file}")
